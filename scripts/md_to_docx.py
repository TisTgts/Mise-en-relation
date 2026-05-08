from pathlib import Path

from docx import Document


def flush_code_block(document: Document, code_lines: list[str]) -> None:
    if not code_lines:
        return
    paragraph = document.add_paragraph("\n".join(code_lines))
    paragraph.style = "No Spacing"
    code_lines.clear()


def convert_markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    document = Document()
    lines = md_path.read_text(encoding="utf-8").splitlines()

    in_code_block = False
    code_lines: list[str] = []

    for raw_line in lines:
        line = raw_line.rstrip()

        if line.strip().startswith("```"):
            in_code_block = not in_code_block
            if not in_code_block:
                flush_code_block(document, code_lines)
            continue

        if in_code_block:
            code_lines.append(raw_line)
            continue

        stripped = line.strip()
        if not stripped:
            document.add_paragraph("")
            continue

        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")), 3)
            title = stripped[level:].strip()
            document.add_heading(title, level=level)
            continue

        if stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
            continue

        if len(stripped) > 2 and stripped[0].isdigit() and stripped[1:3] == ". ":
            document.add_paragraph(stripped[3:].strip(), style="List Number")
            continue

        document.add_paragraph(raw_line)

    flush_code_block(document, code_lines)
    document.save(docx_path)


if __name__ == "__main__":
    project_root = Path(__file__).resolve().parents[1]
    md_file = project_root / "algorithme_matching.md"
    docx_file = project_root / "algorithme_matching.docx"
    convert_markdown_to_docx(md_file, docx_file)
    print(f"Fichier Word cree: {docx_file}")
